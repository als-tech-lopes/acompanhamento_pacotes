import csv
import hashlib
import json
import os
import tempfile
import time
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote
from zoneinfo import ZoneInfo

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import psycopg
import requests
import xlrd
from dateutil import parser as date_parser
from psycopg.types.json import Jsonb

from mapping import map_headers

DATABASE_URL = os.environ["DATABASE_URL"]
SUPABASE_URL = os.environ["SUPABASE_URL"].rstrip("/")
SERVICE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
POLL_SECONDS = int(os.getenv("POLL_INTERVAL_SECONDS", "2"))
BATCH_SIZE = int(os.getenv("BATCH_SIZE", "500"))
STALE_JOB_MINUTES = int(os.getenv("STALE_JOB_MINUTES", "60"))
SOURCE_TIMEZONE = ZoneInfo(os.getenv("SOURCE_TIMEZONE", "America/Sao_Paulo"))
CSV_ENCODING = os.getenv("CSV_ENCODING", "utf-8-sig")
STORAGE_HEADERS = {"Authorization": f"Bearer {SERVICE_KEY}", "apikey": SERVICE_KEY}

SILVER_TEXT_FIELDS = {
    "numero_waybill", "numero_ordem_escaneamento", "numero_pedido_anjun", "tipo_pacote", "status_pacote", "tag", "tipo_pedido",
    "nome_cliente", "nome_cliente_merchant", "ultimo_rastreio", "ultimo_ponto_rastreio", "motivo_ocorrencia", "uf_destinatario",
    "cidade_destinatario", "bairro_destinatario", "rua_destinatario", "cep_destinatario", "regiao", "uf_origem", "cidade_remetente",
    "ponto_pre_entrega", "ponto_entrada", "uf_ponto_entrada", "ponto_entrega", "entregador", "tipo_coleta", "tipo_assinatura"
}
SILVER_DATE_FIELDS = {"data_ultimo_rastreio", "data_criacao_pedido", "data_saida_armazem_cliente", "data_coleta_lm", "data_hora_entrega_ponto", "data_hora_primeira_entrega"}

def clean_value(value):
    if value is None or str(value).strip() in {"", "--", "null", "None"}: return None
    if isinstance(value, (datetime, date)): return value.isoformat()
    return str(value).strip()

def as_datetime(value):
    value = clean_value(value)
    if not value: return None
    try:
        parsed = date_parser.parse(value)
        return parsed.replace(tzinfo=SOURCE_TIMEZONE) if parsed.tzinfo is None else parsed
    except (ValueError, TypeError, OverflowError): return None

def as_number(value, integer=False):
    value = clean_value(value)
    if value is None: return None
    try: return int(float(value.replace(",", "."))) if integer else float(value.replace(",", "."))
    except (ValueError, TypeError): return None

def claim_job(connection):
    with connection.transaction():
        return connection.execute("""
          update public.importacoes set status='processando', etapa='download', progresso=10, mensagem_erro=null, updated_at=now()
          where id = (select id from public.importacoes where status='aguardando_processamento' order by iniciada_em for update skip locked limit 1)
          returning id, nome_arquivo, storage_original_path, tamanho_bytes
        """).fetchone()

def recover_stale_jobs(connection):
    stale_ids = connection.execute("""
      select id from public.importacoes
      where status='processando'
        and updated_at < now() - make_interval(mins => %s)
      for update skip locked
    """, (STALE_JOB_MINUTES,)).fetchall()
    for (import_id,) in stale_ids:
        connection.execute("delete from public.registros_processados where importacao_id=%s", (import_id,))
        connection.execute("delete from public.bronze_entregas_raw where importacao_id=%s", (import_id,))
        connection.execute("delete from public.arquivos_processados where importacao_id=%s", (import_id,))
        connection.execute("""
          update public.importacoes
          set status='aguardando_processamento', etapa='recuperacao', progresso=5,
              mensagem_erro='Carga recuperada após interrupção do worker.', updated_at=now()
          where id=%s
        """, (import_id,))
    if stale_ids:
        connection.commit()
        print(f"{len(stale_ids)} carga(s) interrompida(s) devolvida(s) à fila", flush=True)

def update_job(connection, import_id, **values):
    assignments = ", ".join(f"{key}=%s" for key in values)
    connection.execute(f"update public.importacoes set {assignments}, updated_at=now() where id=%s", [*values.values(), import_id])
    connection.commit()

def download_file(storage_path, suffix):
    response = requests.get(f"{SUPABASE_URL}/storage/v1/object/importacoes/{quote(storage_path, safe='/')}", headers=STORAGE_HEADERS, stream=True, timeout=120)
    response.raise_for_status()
    temp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    with temp:
        for chunk in response.iter_content(1024 * 1024):
            if chunk: temp.write(chunk)
    return temp.name

def hash_file(file_path):
    digest = hashlib.sha256()
    with open(file_path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def delete_original(storage_path):
    try:
        requests.delete(f"{SUPABASE_URL}/storage/v1/object/importacoes/{quote(storage_path, safe='/')}", headers=STORAGE_HEADERS, timeout=30).raise_for_status()
    except requests.RequestException as error:
        print(f"Não foi possível remover o arquivo duplicado do Storage: {error}", flush=True)

def iter_spreadsheet(file_path):
    extension = Path(file_path).suffix.lower()
    if extension == ".xlsx":
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try: headers = [clean_value(v) or "" for v in next(rows)]
        except StopIteration: raise ValueError("Planilha vazia.")
        return sheet.title, headers, rows, workbook.close
    if extension == ".xls":
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        headers = [clean_value(v) or "" for v in sheet.row_values(0)]
        return sheet.name, headers, (sheet.row_values(i) for i in range(1, sheet.nrows)), workbook.release_resources
    raw = open(file_path, "r", encoding=CSV_ENCODING, newline="")
    sample = raw.read(8192); raw.seek(0)
    try: dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error: dialect = csv.excel
    rows = csv.reader(raw, dialect)
    try: headers = [clean_value(v) or "" for v in next(rows)]
    except StopIteration: raise ValueError("Arquivo CSV vazio.")
    return Path(file_path).stem, headers, rows, raw.close

def make_payload(headers, values):
    padded = list(values) + [None] * max(0, len(headers) - len(values))
    return {"columns": [{"position": index + 1, "header": header, "value": clean_value(padded[index])} for index, header in enumerate(headers)]}

def mapped_row(mapping, values):
    padded = list(values) + [None] * max(0, len(mapping) - len(values))
    return {item["technical"]: clean_value(padded[index]) for index, item in enumerate(mapping) if item["technical"]}

def silver_values(mapped):
    values = {field: mapped.get(field) for field in SILVER_TEXT_FIELDS}
    values.update({field: as_datetime(mapped.get(field)) for field in SILVER_DATE_FIELDS})
    values["duracao_atraso_horas"] = as_number(mapped.get("duracao_atraso_horas"))
    values["dias_atraso"] = as_number(mapped.get("dias_atraso"), integer=True)
    values["prazo_calculado"] = as_datetime(mapped.get("horario_limite_entrega"))
    latest = values.get("data_ultimo_rastreio")
    if latest:
        if latest.tzinfo is None: latest = latest.replace(tzinfo=timezone.utc)
        values["horas_sem_movimentacao"] = max(0, (datetime.now(timezone.utc) - latest).total_seconds() / 3600)
    else: values["horas_sem_movimentacao"] = None
    values["status_sla"] = "fora_prazo" if (values.get("dias_atraso") or 0) > 0 else "no_prazo"
    return values

def flush_batch(connection, import_id, file_name, headers, mapping, batch):
    for line_number, values, record_hash in batch:
        payload = make_payload(headers, values)
        mapped = mapped_row(mapping, values)
        errors = [field for field in ("numero_waybill", "status_pacote", "data_criacao_pedido") if not mapped.get(field)]
        bronze_id = connection.execute("""
          insert into public.bronze_entregas_raw
          (importacao_id, numero_linha_origem, payload_original_json, cabecalhos_originais, nome_arquivo_origem, hash_registro, classificacao_validacao, erros_validacao)
          values (%s,%s,%s,%s,%s,%s,%s,%s) returning id
        """, (import_id, line_number, Jsonb(payload), Jsonb(headers), file_name, record_hash, "erro" if errors else "valido", Jsonb(errors))).fetchone()[0]
        silver = silver_values(mapped)
        columns = ["bronze_id", "importacao_id", *silver.keys(), "possui_erro", "erros_transformacao", "valores_originais_criticos"]
        placeholders = ",".join(["%s"] * len(columns))
        connection.execute(f"insert into public.silver_entregas ({','.join(columns)}) values ({placeholders})", [bronze_id, import_id, *silver.values(), bool(errors), Jsonb(errors), Jsonb({})])
    connection.commit()

def upload_parquet(file_path, storage_path):
    with open(file_path, "rb") as body:
        response = requests.post(f"{SUPABASE_URL}/storage/v1/object/parquet/{quote(storage_path, safe='/')}", headers={**STORAGE_HEADERS, "Content-Type": "application/vnd.apache.parquet", "x-upsert": "true"}, data=body, timeout=300)
    response.raise_for_status()

def process_job(connection, job):
    import_id, file_name, storage_path, _ = job
    suffix = Path(file_name).suffix.lower()
    source_path = parquet_path = file_hash = None
    try:
        source_path = download_file(storage_path, suffix)
        file_hash = hash_file(source_path)
        claimed_file = connection.execute("""
          insert into public.arquivos_processados (hash_arquivo, importacao_id)
          values (%s, %s) on conflict (hash_arquivo) do nothing returning hash_arquivo
        """, (file_hash, import_id)).fetchone()
        if not claimed_file:
            original = connection.execute("select importacao_id from public.arquivos_processados where hash_arquivo=%s", (file_hash,)).fetchone()
            connection.commit()
            delete_original(storage_path)
            update_job(connection, import_id, status="duplicada", etapa="deduplicacao", progresso=100, hash_arquivo=file_hash, total_linhas=0, linhas_processadas=0, linhas_duplicadas=0, concluida_em=datetime.now(timezone.utc), metadados=Jsonb({"importacao_original_id": str(original[0]) if original else None, "motivo": "arquivo_identico"}))
            print(f"Importação {import_id} ignorada: arquivo idêntico já processado", flush=True)
            return
        connection.commit()
        update_job(connection, import_id, hash_arquivo=file_hash, etapa="leitura", progresso=20)
        sheet_name, headers, rows, close_reader = iter_spreadsheet(source_path)
        mapping = map_headers(headers)
        required = {"numero_waybill", "status_pacote", "data_criacao_pedido"}
        available = {item["technical"] for item in mapping if item["technical"]}
        missing = sorted(required - available)
        if missing: raise ValueError(f"Colunas obrigatórias ausentes: {', '.join(missing)}")
        schema = pa.schema([pa.field(f"col_{index:03d}", pa.string()) for index in range(1, len(headers) + 1)], metadata={b"original_headers": json.dumps(headers, ensure_ascii=False).encode("utf-8")})
        parquet_file = tempfile.NamedTemporaryFile(delete=False, suffix=".parquet"); parquet_path = parquet_file.name; parquet_file.close()
        writer = pq.ParquetWriter(parquet_path, schema, compression="zstd")
        batch = []; parquet_rows = []; total = valid = invalid = duplicates = 0
        update_job(connection, import_id, etapa="bronze_silver", progresso=30, total_colunas=len(headers), colunas_mapeadas=len(available), colunas_nao_mapeadas=sum(1 for item in mapping if not item["mapped"]), metadados=Jsonb({"sheetName": sheet_name, "sheetCount": 1, "mapping": mapping}))
        try:
            for line_number, values in enumerate(rows, 2):
                if not any(clean_value(v) is not None for v in values): continue
                total += 1
                payload = make_payload(headers, values)
                serialized = json.dumps(payload, ensure_ascii=False, default=str, separators=(",", ":"))
                record_hash = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
                claimed_record = connection.execute("""
                  insert into public.registros_processados (hash_registro, importacao_id)
                  values (%s, %s) on conflict (hash_registro) do nothing returning hash_registro
                """, (record_hash, import_id)).fetchone()
                if not claimed_record:
                    duplicates += 1
                    continue
                mapped = mapped_row(mapping, values)
                errors = [field for field in required if not mapped.get(field)]
                invalid += bool(errors); valid += not errors
                batch.append((line_number, values, record_hash))
                padded = list(values) + [None] * max(0, len(headers) - len(values))
                parquet_rows.append({f"col_{i+1:03d}": clean_value(padded[i]) for i in range(len(headers))})
                if len(batch) >= BATCH_SIZE:
                    flush_batch(connection, import_id, file_name, headers, mapping, batch)
                    writer.write_table(pa.Table.from_pylist(parquet_rows, schema=schema))
                    batch.clear(); parquet_rows.clear()
                    update_job(connection, import_id, linhas_processadas=total, total_linhas=total, linhas_validas=valid, linhas_erro=invalid, progresso=min(85, 30 + total // BATCH_SIZE))
            if batch:
                flush_batch(connection, import_id, file_name, headers, mapping, batch)
                writer.write_table(pa.Table.from_pylist(parquet_rows, schema=schema))
        finally:
            writer.close(); close_reader()
        connection.execute("""
          with ranked as (select id, row_number() over(partition by numero_waybill order by id) rn from public.silver_entregas where importacao_id=%s and numero_waybill is not null)
          update public.silver_entregas s set eh_duplicado=true from ranked r where s.id=r.id and r.rn>1
        """, (import_id,))
        duplicates += connection.execute("select count(*) from public.silver_entregas where importacao_id=%s and eh_duplicado", (import_id,)).fetchone()[0]
        connection.commit()
        update_job(connection, import_id, etapa="parquet", progresso=90, total_linhas=total, linhas_processadas=total, linhas_validas=valid, linhas_erro=invalid, linhas_duplicadas=duplicates)
        parquet_storage_path = f"{import_id}/bronze.parquet"; upload_parquet(parquet_path, parquet_storage_path)
        update_job(connection, import_id, status="concluida", etapa="gold", progresso=100, storage_parquet_path=parquet_storage_path, concluida_em=datetime.now(timezone.utc))
        print(f"Importação {import_id} concluída: {valid + invalid} novas, {duplicates} duplicadas", flush=True)
    except Exception as error:
        connection.rollback()
        connection.execute("delete from public.registros_processados where importacao_id=%s", (import_id,))
        connection.execute("delete from public.bronze_entregas_raw where importacao_id=%s", (import_id,))
        connection.execute("delete from public.arquivos_processados where importacao_id=%s", (import_id,))
        connection.commit()
        update_job(connection, import_id, status="falhou", etapa="erro", mensagem_erro=str(error)[:2000])
        print(f"Importação {import_id} falhou: {error}", flush=True)
    finally:
        for path in (source_path, parquet_path):
            if path:
                try: os.unlink(path)
                except OSError: pass

def main():
    print("Worker de importações iniciado", flush=True)
    while True:
        try:
            with psycopg.connect(DATABASE_URL) as connection:
                recover_stale_jobs(connection)
                job = claim_job(connection)
                if job: process_job(connection, job)
        except Exception as error:
            print(f"Erro no ciclo do worker: {error}", flush=True)
        time.sleep(POLL_SECONDS)

if __name__ == "__main__": main()
